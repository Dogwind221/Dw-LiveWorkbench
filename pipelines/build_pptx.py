#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
模块 3 渲染器：把 DeckSpec(JSON) 渲染成真正的 .pptx

用法：
    python build_pptx.py <deck.json> <out.pptx>

设计说明（2026-10 重做）：
  - 16:9，13.333 x 7.5 英寸
  - 去掉「色块卡片 + 彩色圆点」那套模板感：改用编辑设计（editorial）的版面语言
    —— 细分隔线、幽灵序号、克制的字距、大留白
  - 茶叶叶形轮廓用 freeform 参数化绘制（见 add_leaf），做封面/深色页的品牌纹样
  - 每张幻灯片的演讲者备注 = 该页口播稿
"""
import json
import math
import sys
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- 设计系统
INK = "14302A"       # 深墨绿：封面与深色页的底
INK_2 = "1D3E36"     # 渐变末端
PRIMARY = "2C5A4C"   # 主绿
PRIMARY_L = "4C8672"
TINT = "EAF1EC"      # 浅绿底
CREAM = "FAF7F2"     # 页面底色
PAPER = "FFFFFF"
TEXT = "23302B"
MUTED = "7C8A84"
HAIRLINE = "E7E0D4"  # 暖灰分隔线
GHOST = "E4DCCB"     # 幽灵序号
AMBER = "B98A3C"     # 茶金
AMBER_L = "E2C48C"
ON_DARK = "C9D8D2"
ON_DARK_MUTED = "8FA79E"

FONT = "微软雅黑"
SW, SH = 13.333, 7.5
ML, MR = 0.95, 0.95
CW = SW - ML - MR                    # 11.433

KICKER_Y = 0.50
TITLE_Y = 0.80
RULE_Y = 1.58
CONTENT_TOP = 2.00
CONTENT_BOTTOM = 6.56
FOOT_LINE = 6.86
FOOTER_Y = 6.98

EMU_PER_IN = 914400


def rgb(h):
    return RGBColor.from_string(h)


def set_font(run, name=FONT, size=None, bold=None, color=None, italic=None):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = rgb(color)
    f.name = name
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {"typeface": name})
            if latin is not None:
                latin.addnext(el)
            else:
                rPr.append(el)
        else:
            el.set("typeface", name)
    return run


def set_tracking(run, hundredths):
    """字距：a:rPr/@spc，单位 1/100 磅。小字号加一点字距会明显更精致。"""
    run._r.get_or_add_rPr().set("spc", str(int(hundredths)))
    return run


def set_strike(run):
    """strike 是 a:rPr 的属性（不是子元素）。"""
    run._r.get_or_add_rPr().set("strike", "sngStrike")
    return run


def set_alpha(shape, pct):
    """给纯色填充加透明度（python-pptx 没有 API，只能改 XML）。pct = 不透明度百分比。"""
    try:
        spPr = shape._element.spPr
        solid = spPr.find(qn("a:solidFill"))
        if solid is None:
            return False
        clr = solid.find(qn("a:srgbClr"))
        if clr is None:
            return False
        for old in clr.findall(qn("a:alpha")):
            clr.remove(old)
        clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))}))
        return True
    except Exception:
        return False


def no_autofit(tf):
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    if bodyPr is not None:
        for tag in ("a:normAutofit", "a:spAutoFit"):
            el = bodyPr.find(qn(tag))
            if el is not None:
                bodyPr.remove(el)
    return tf


def rect(slide, x, y, w, h, fill=None, line=None, radius=None, line_w=0.75, shape=None):
    shp = slide.shapes.add_shape(
        shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.word_wrap = True
    return shp


def oval(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    return rect(slide, x, y, w, h, fill=fill, line=line, line_w=line_w, shape=MSO_SHAPE.OVAL)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    no_autofit(tf)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def para(tf, text, size=14, bold=False, color=TEXT, align=PP_ALIGN.LEFT, first=False,
         space_before=0, space_after=0, line_spacing=1.25, font=FONT, tracking=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if space_before:
        p.space_before = Pt(space_before)
    if space_after:
        p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = "" if text is None else str(text)
    set_font(r, name=font, size=size, bold=bold, color=color)
    if tracking:
        set_tracking(r, tracking)
    return p


def shape_text(shp, text, size=11, bold=True, color=INK, align=PP_ALIGN.CENTER,
               tracking=None, pad=0.0):
    """把文字写进一个已有形状（胶囊标签 / 表头用）。"""
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(pad)
    tf.margin_right = Inches(pad)
    tf.margin_top = 0
    tf.margin_bottom = 0
    para(tf, text, size=size, bold=bold, color=color, align=align, first=True, tracking=tracking)
    return shp


def cjk_lines(text, width_in, size_pt):
    """粗略估算中文段落会折几行（每字宽 ≈ 字号）。用于排版时判断内容块高度。"""
    per_line = max(1, int(width_in / max(0.01, size_pt / 72.0)))
    return max(1, math.ceil(len(str(text)) / per_line))


def block_height(text, width_in, size_pt, line_spacing=1.4):
    return cjk_lines(text, width_in, size_pt) * (size_pt * line_spacing / 72.0)


def centered_offset(available, needed):
    """内容比版面矮时返回居中偏移量，否则返回 0。

    卡片/清单类版面最容易犯的毛病是内容顶在上沿、下面空一大片，
    看起来像没排完。按内容实高居中就整齐了。
    """
    return max(0.0, (available - needed) / 2.0)


def leaf_points(cx, cy, length, width, angle_deg, n=44):
    """参数化茶叶轮廓：上缘饱满、下缘收成尖基，再绕中心旋转。"""
    pts = []
    for i in range(n + 1):
        t = i / n
        x = (t - 0.5) * length
        y = -width * (math.sin(math.pi * t) ** 1.05)
        pts.append((x, y))
    for i in range(n, -1, -1):
        t = i / n
        x = (t - 0.5) * length
        y = width * 0.8 * (math.sin(math.pi * t) ** 1.6)
        pts.append((x, y))
    a = math.radians(angle_deg)
    return [(cx + x * math.cos(a) - y * math.sin(a), cy + x * math.sin(a) + y * math.cos(a))
            for x, y in pts]


def add_leaf(slide, cx, cy, length, width, angle, fill, alpha=100):
    """画一片叶子（品牌纹样，允许出血到画布外）。

    命名以 `deco ` 开头：版式检查器据此把「刻意的装饰出血」和「内容溢出」区分开，
    不用靠猜页码或猜有没有文字。

    局部坐标必须用 EMU（scale=1）：python-pptx 把 path 坐标按整数写入，
    若局部单位是英寸，坐标会被取整成 0/1/2，轮廓直接碎成阶梯状。
    """
    pts = [(x * EMU_PER_IN, y * EMU_PER_IN) for x, y in leaf_points(cx, cy, length, width, angle)]
    builder = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    builder.add_line_segments(pts[1:], close=True)
    shp = builder.convert_to_shape()
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if alpha < 100:
        set_alpha(shp, alpha)
    shp.line.fill.background()
    shp.shadow.inherit = False
    shp.name = "deco leaf"
    return shp


def mark_deco(shp):
    """把形状标记为装饰（允许出血）。"""
    shp.name = "deco " + (shp.name or "shape")
    return shp


# ---------------------------------------------------------------- 页面零件
def page_base(slide, dark=False):
    rect(slide, 0, 0, SW, SH, fill=INK if dark else CREAM)


def header(slide, kicker, title, subtitle=""):
    """页眉：金色圆点 + 拉开字距的小标签，标题，然后「金色短段 + 整条细线」的分隔。

    没有 kicker 时（例如开场页）标题位置上移，不留空档。
    """
    title_y = TITLE_Y
    if kicker:
        oval(slide, ML, KICKER_Y + 0.075, 0.075, 0.075, fill=AMBER)
        tf = textbox(slide, ML + 0.175, KICKER_Y, CW - 0.2, 0.28)
        para(tf, kicker, size=11, bold=True, color=AMBER, first=True, tracking=140)
    else:
        title_y = TITLE_Y - 0.24

    tf = textbox(slide, ML, title_y, CW, 0.72)
    para(tf, title, size=31, bold=True, color=INK, first=True, line_spacing=1.08)

    rule_y = RULE_Y - (0 if kicker else 0.24)
    rect(slide, ML, rule_y, CW, 0.011, fill=HAIRLINE)
    rect(slide, ML, rule_y - 0.012, 0.95, 0.035, fill=AMBER)

    top = rule_y + (CONTENT_TOP - RULE_Y)
    if subtitle:
        tf = textbox(slide, ML, rule_y + 0.17, CW, 0.32)
        para(tf, subtitle, size=12.5, color=MUTED, first=True)
        top += 0.30
    return top


def footer(slide, left, right="", dark=False):
    rect(slide, ML, FOOT_LINE, CW, 0.008, fill="2C4A42" if dark else HAIRLINE)
    if left:
        tf = textbox(slide, ML, FOOTER_Y, CW * 0.75, 0.26)
        para(tf, left, size=9, color=ON_DARK_MUTED if dark else MUTED, first=True, tracking=60)
    if right:
        tf = textbox(slide, SW - MR - 2.6, FOOTER_Y, 2.6, 0.26)
        para(tf, right, size=9, color=ON_DARK_MUTED if dark else MUTED, align=PP_ALIGN.RIGHT,
             first=True, tracking=60)


def notes(slide, deck, s):
    parts = []
    if s.get("script"):
        parts.append("【口播稿】")
        for line in s["script"]:
            cue = line.get("cue")
            parts.append(("〔%s〕" % cue if cue else "") + line.get("text", ""))
    slide.notes_slide.notes_text_frame.text = "\n".join(parts)


def foot_left(deck, s):
    brand = deck.get("meta", {}).get("brand", "")
    sec = s.get("sectionTitle", "")
    return " · ".join([x for x in (brand, sec) if x])


# ---------------------------------------------------------------- 封面
def slide_cover(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = rect(slide, 0, 0, SW, SH, fill=INK)
    try:  # 竖向渐变，比纯色更有体积
        fill = bg.fill
        fill.gradient()
        fill.gradient_angle = 90.0
        stops = fill.gradient_stops
        stops[0].color.rgb = rgb(INK)
        stops[0].position = 0.0
        stops[1].color.rgb = rgb(INK_2)
        stops[1].position = 1.0
    except Exception:
        bg.fill.solid()
        bg.fill.fore_color.rgb = rgb(INK)
    bg.line.fill.background()

    add_leaf(slide, 10.9, 2.45, 5.6, 2.0, -22, "FFFFFF", alpha=7)
    add_leaf(slide, 12.1, 5.85, 3.4, 1.15, 14, AMBER_L, alpha=10)
    rect(slide, 0, 0, 0.14, SH, fill=AMBER)

    mark_deco(oval(slide, 8.95, 0.95, 4.5, 4.5, fill=None, line="254A40", line_w=1.0))
    mark_deco(oval(slide, 9.85, 1.85, 2.7, 2.7, fill=None, line="2E564B", line_w=1.0))
    oval(slide, 11.00, 3.00, 0.42, 0.42, fill=AMBER)

    brand = deck.get("meta", {}).get("brand", "")
    if brand:
        tf = textbox(slide, SW - MR - 3.0, 0.62, 3.0, 0.3)
        para(tf, brand, size=10, color="7E958C", align=PP_ALIGN.RIGHT, first=True, tracking=120)

    if s.get("kicker"):
        pill = rect(slide, ML, 1.36, max(1.5, 0.30 + len(s["kicker"]) * 0.20), 0.38,
                    fill=AMBER, radius=0.5)
        shape_text(pill, s["kicker"], size=11, bold=True, color=INK, tracking=100)

    tf = textbox(slide, ML, 1.98, 9.2, 2.2)
    para(tf, s.get("title", ""), size=44, bold=True, color=CREAM, first=True, line_spacing=1.06)

    rect(slide, ML, 4.42, 1.15, 0.045, fill=AMBER)

    if s.get("subtitle"):
        tf = textbox(slide, ML, 4.72, 9.0, 0.55)
        para(tf, s["subtitle"], size=18, color=ON_DARK, first=True, line_spacing=1.25)

    rect(slide, ML, 6.06, 9.4, 0.008, fill="3A5A51")
    meta = s.get("meta") or []
    if meta:
        tf = textbox(slide, ML, 6.24, 9.0, 0.9)
        for i, m in enumerate(meta):
            para(tf, m, size=11, color=ON_DARK_MUTED, first=(i == 0), space_after=6, line_spacing=1.4)

    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 要点
def slide_points(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    top = header(slide, s.get("kicker"), s.get("title"), s.get("subtitle", ""))

    items = s.get("bullets") or []
    quote = s.get("quote")
    bottom = CONTENT_BOTTOM - (0.96 if quote else 0)
    n = max(1, len(items))
    avail = bottom - top
    row_h = min(1.12, avail / n)
    start = top + centered_offset(avail, row_h * n)

    for i, it in enumerate(items):
        cy = start + i * row_h + row_h / 2
        tf = textbox(slide, ML - 0.05, cy - 0.3, 0.8, 0.6, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, "%02d" % (i + 1), size=22, bold=True, color=GHOST, align=PP_ALIGN.RIGHT, first=True)
        tf = textbox(slide, ML + 1.02, cy - 0.42, CW - 1.12, 0.84, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, it, size=16, color=TEXT, first=True, line_spacing=1.28)
        if i < n - 1:
            rect(slide, ML, start + (i + 1) * row_h - 0.01, CW, 0.007, fill=HAIRLINE)

    if quote:
        qy = CONTENT_BOTTOM - 0.80
        rect(slide, ML, qy, CW, 0.80, fill=INK, radius=0.06)
        rect(slide, ML, qy, 0.055, 0.80, fill=AMBER)
        tf = textbox(slide, ML + 0.42, qy, CW - 0.84, 0.80, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, quote, size=15.5, bold=True, color=AMBER_L, first=True, line_spacing=1.2)

    footer(slide, foot_left(deck, s), "%02d" % s["index"])
    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 卡片
def slide_cards(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    top = header(slide, s.get("kicker"), s.get("title"), s.get("subtitle", ""))

    cards = s.get("cards") or []
    quote = s.get("quote")
    bottom = CONTENT_BOTTOM - (1.06 if quote else 0)
    n = len(cards) or 1
    if n <= 3:
        cols, rows = n, 1
    else:
        cols, rows = 2, (n + 1) // 2
    gap = 0.30
    cw = (CW - gap * (cols - 1)) / cols
    ch = (bottom - top - gap * (rows - 1)) / rows

    for i, c in enumerate(cards):
        r_, c_ = divmod(i, cols)
        x = ML + c_ * (cw + gap)
        y = top + r_ * (ch + gap)
        rect(slide, x, y, cw, ch, fill=PAPER, line=HAIRLINE, line_w=0.75, radius=0.06)

        # 卡片内容按实高居中：短文案不再缩在卡片上沿
        inner_w = cw - 0.72
        desc = c.get("desc", "")
        desc_h = block_height(desc, inner_w, 11.5, 1.42)
        needed = 0.20 + 0.42 + 0.16 + desc_h + 0.36
        offset = centered_offset(ch, needed)
        iy = y + 0.36 + offset

        rect(slide, x + 0.36, iy, 0.34, 0.032, fill=AMBER)
        tf = textbox(slide, x + cw - 0.86, iy - 0.12, 0.6, 0.5)
        para(tf, "%02d" % (i + 1), size=19, bold=True, color=GHOST, align=PP_ALIGN.RIGHT, first=True)
        tf = textbox(slide, x + 0.36, iy + 0.20, inner_w, 0.42)
        para(tf, c.get("title", ""), size=15, bold=True, color=PRIMARY, first=True)
        # 描述框高度按「卡片下沿留 0.36 边距」反算，不能再用固定值：
        # 内容居中后起点下移了，固定高度会顶出卡片外（版式检查器会报越界）
        desc_box_h = max(0.30, (y + ch - 0.36) - (iy + 0.68))
        tf = textbox(slide, x + 0.36, iy + 0.68, inner_w, desc_box_h)
        para(tf, desc, size=11.5, color=TEXT, first=True, line_spacing=1.42)

    if quote:
        qy = CONTENT_BOTTOM - 0.88
        rect(slide, ML, qy, CW, 0.88, fill=TINT, radius=0.06)
        rect(slide, ML, qy, 0.055, 0.88, fill=PRIMARY_L)
        tf = textbox(slide, ML + 0.42, qy, CW - 0.84, 0.88, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, quote, size=13.5, color=PRIMARY, first=True, line_spacing=1.32)

    footer(slide, foot_left(deck, s), "%02d" % s["index"])
    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 清单（原「表格」）
def slide_table(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    top = header(slide, s.get("kicker"), s.get("title"), s.get("subtitle", ""))

    rows = (s.get("table") or {}).get("rows") or []
    n = max(1, len(rows))
    avail = CONTENT_BOTTOM - top
    row_h = min(0.88, avail / n)
    start = top + centered_offset(avail, row_h * n)
    lab_w = 2.25

    # 标签列与内容列之间一条细分隔线，比色块表格克制得多
    rect(slide, ML + lab_w - 0.30, start + 0.06, 0.008, max(0.2, row_h * n - 0.12), fill=HAIRLINE)

    for i, row in enumerate(rows):
        y = start + i * row_h
        vals = list(row) + ["", ""]
        label, value = vals[0], vals[1]
        tf = textbox(slide, ML, y, lab_w - 0.48, row_h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, label, size=10.5, color=MUTED, first=True, tracking=110)
        tf = textbox(slide, ML + lab_w, y, CW - lab_w, row_h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, value, size=14.5, color=TEXT, first=True, line_spacing=1.25)
        if i < n - 1:
            rect(slide, ML, y + row_h - 0.01, CW, 0.007, fill=HAIRLINE)

    if s.get("footnote"):
        tf = textbox(slide, ML, CONTENT_BOTTOM - 0.06, CW, 0.3)
        para(tf, s["footnote"], size=10.5, color=MUTED, first=True)

    footer(slide, foot_left(deck, s), "%02d" % s["index"])
    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 对比
def slide_compare(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    top = header(slide, s.get("kicker"), s.get("title"), s.get("subtitle", ""))

    gap = 0.36
    cw = (CW - gap) / 2
    ch = CONTENT_BOTTOM - top

    for i, side in enumerate([s.get("left"), s.get("right")]):
        if not side:
            continue
        x = ML + i * (cw + gap)
        good = i == 0
        rect(slide, x, top, cw, ch, fill=PAPER, line=HAIRLINE, line_w=0.75)
        head = rect(slide, x, top, cw, 0.60, fill=INK if good else "8E9A94")
        shape_text(head, side.get("title", ""), size=15, bold=True, color=PAPER,
                   align=PP_ALIGN.LEFT, pad=0.34)
        items = side.get("items") or []
        ih = min(0.88, (ch - 0.86) / max(1, len(items)))
        for j, it in enumerate(items):
            y = top + 0.84 + j * ih
            oval(slide, x + 0.36, y + 0.13, 0.11, 0.11, fill=AMBER if good else "B7C2BD")
            tf = textbox(slide, x + 0.64, y, cw - 1.0, ih)
            para(tf, it, size=12.5, color=TEXT, first=True, line_spacing=1.38)

    footer(slide, foot_left(deck, s), "%02d" % s["index"])
    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 步骤
def slide_steps(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    top = header(slide, s.get("kicker"), s.get("title"), s.get("subtitle", ""))

    steps = s.get("steps") or []
    n = max(1, len(steps))
    dense = n >= 5
    r = 0.27 if dense else 0.32
    gap = 0.14
    cw = (CW - gap * (n - 1)) / n

    # 步骤块整体在正文区垂直居中：步骤少的时候不再贴着页眉、下面空一大截
    desc_size = 10.5 if dense else 11
    max_desc_h = max([block_height(st.get("desc", ""), cw - 0.1, desc_size, 1.36)
                      for st in steps] or [0.4])
    need = r * 2 + 0.22 + 0.38 + max_desc_h + 0.34
    avail = (CONTENT_BOTTOM - (0.86 if s.get("footnote") else 0)) - top
    cy = top + centered_offset(avail, need) + r

    rect(slide, ML + cw / 2, cy - 0.005, CW - cw, 0.008, fill=HAIRLINE)

    for i, st in enumerate(steps):
        x = ML + i * (cw + gap)
        cx = x + cw / 2
        if i == 0:
            oval(slide, cx - r, cy - r, r * 2, r * 2, fill=PRIMARY)
        else:
            oval(slide, cx - r, cy - r, r * 2, r * 2, fill=PAPER, line=HAIRLINE, line_w=1.0)
        tf = textbox(slide, cx - 0.3, cy - 0.22, 0.6, 0.44, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, str(i + 1), size=(14 if dense else 16), bold=True,
             color=PAPER if i == 0 else PRIMARY, align=PP_ALIGN.CENTER, first=True)
        tf = textbox(slide, x, cy + r + 0.22, cw, 0.38)
        para(tf, st.get("title", ""), size=(13 if dense else 14), bold=True, color=PRIMARY,
             align=PP_ALIGN.CENTER, first=True)
        tf = textbox(slide, x + 0.05, cy + r + 0.64, cw - 0.1, 1.5)
        para(tf, st.get("desc", ""), size=(10.5 if dense else 11), color=TEXT,
             align=PP_ALIGN.CENTER, first=True, line_spacing=1.36)

    if s.get("footnote"):
        rect(slide, ML, CONTENT_BOTTOM - 0.72, CW, 0.64, fill=TINT, radius=0.10)
        rect(slide, ML, CONTENT_BOTTOM - 0.72, 0.055, 0.64, fill=PRIMARY_L)
        tf = textbox(slide, ML + 0.4, CONTENT_BOTTOM - 0.72, CW - 0.8, 0.64, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, s["footnote"], size=12.5, color=PRIMARY, first=True)

    footer(slide, foot_left(deck, s), "%02d" % s["index"])
    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 答疑
def slide_faq(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    top = header(slide, s.get("kicker"), s.get("title"), s.get("subtitle", ""))

    items = s.get("faq") or []
    n = max(1, len(items))
    avail = CONTENT_BOTTOM - top
    row_h = min(1.12, avail / n)
    start = top + centered_offset(avail, row_h * n)

    for i, f in enumerate(items):
        y = start + i * row_h
        tag = rect(slide, ML, y + 0.03, 0.30, 0.30, fill=AMBER, radius=0.24)
        shape_text(tag, "Q", size=10.5, bold=True, color=PAPER)
        tf = textbox(slide, ML + 0.46, y, CW - 0.46, 0.34)
        para(tf, f.get("q", ""), size=13.5, bold=True, color=PRIMARY, first=True)
        tf = textbox(slide, ML + 0.46, y + 0.36, CW - 0.6, max(0.3, row_h - 0.44))
        para(tf, f.get("a", ""), size=11.5, color=TEXT, first=True, line_spacing=1.42)
        if i < n - 1:
            rect(slide, ML, y + row_h - 0.02, CW, 0.007, fill=HAIRLINE)

    footer(slide, foot_left(deck, s), "%02d" % s["index"])
    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 价格
def slide_price(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    top = header(slide, s.get("kicker"), s.get("title"), s.get("subtitle", ""))

    lw = 4.85
    rect(slide, ML, top, lw, 3.05, fill=INK, radius=0.05)
    rect(slide, ML, top, 0.055, 3.05, fill=AMBER)
    tf = textbox(slide, ML + 0.44, top + 0.36, lw - 0.88, 0.3)
    para(tf, "直播间价格", size=10, color=AMBER, first=True, tracking=140)
    if s.get("price"):
        tf = textbox(slide, ML + 0.44, top + 0.72, lw - 0.88, 1.15, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, s["price"], size=50, bold=True, color=AMBER_L, first=True)
    if s.get("comparePrice"):
        tf = textbox(slide, ML + 0.44, top + 2.02, lw - 0.88, 0.42)
        p = para(tf, "日常价 " + s["comparePrice"], size=14, color="93A79E", first=True)
        set_strike(p.runs[0])
    tf = textbox(slide, ML + 0.44, top + 2.52, lw - 0.88, 0.42)
    para(tf, "直播期间有效，结束后恢复日常价", size=10.5, color="7E958C", first=True)

    rx = ML + lw + 0.32
    rw = CW - lw - 0.32
    yy = top
    for a in (s.get("anchors") or [])[:3]:
        rect(slide, rx, yy, rw, 0.84, fill=PAPER, line=HAIRLINE, line_w=0.75, radius=0.05)
        tf = textbox(slide, rx + 0.34, yy, rw - 0.68, 0.84, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, a.get("label", ""), size=9.5, color=MUTED, first=True, tracking=110, space_after=2)
        para(tf, a.get("value", ""), size=19, bold=True, color=PRIMARY)
        yy += 0.95
    for b in (s.get("bullets") or []):
        oval(slide, rx + 0.02, yy + 0.12, 0.09, 0.09, fill=AMBER)
        tf = textbox(slide, rx + 0.24, yy, rw - 0.3, 0.44)
        para(tf, b, size=12, color=TEXT, first=True, line_spacing=1.3)
        yy += 0.42

    footer(slide, foot_left(deck, s), "%02d" % s["index"])
    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 产品
def slide_product(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    top = header(slide, s.get("kicker"), s.get("title"), s.get("subtitle", ""))

    gap = 0.32
    lw = CW * 0.51
    rw = CW - lw - gap
    ch = CONTENT_BOTTOM - top

    rect(slide, ML, top, lw, ch, fill=PAPER, line=HAIRLINE, line_w=0.75, radius=0.06)
    rect(slide, ML + 0.36, top + 0.36, 0.34, 0.032, fill=AMBER)
    tf = textbox(slide, ML + 0.36, top + 0.52, lw - 0.72, 0.32)
    para(tf, "这款茶值得记住的地方", size=10.5, color=AMBER, first=True, tracking=120)
    ly = top + 0.98
    for it in (s.get("bullets") or [])[:4]:
        oval(slide, ML + 0.38, ly + 0.12, 0.10, 0.10, fill=AMBER)
        tf = textbox(slide, ML + 0.66, ly, lw - 1.02, 0.9)
        para(tf, it, size=13.5, color=TEXT, first=True, line_spacing=1.4)
        ly += 0.86

    rx = ML + lw + gap
    rect(slide, rx, top, rw, ch, fill=TINT, radius=0.06)
    rows = (s.get("table") or {}).get("rows") or []
    n = max(1, len(rows))
    rh = min(0.74, (ch - 0.6) / n)
    ry = top + 0.34
    for i, row in enumerate(rows):
        vals = list(row) + ["", ""]
        label, value = vals[0], vals[1]
        tf = textbox(slide, rx + 0.34, ry, 1.25, rh, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, label, size=10, color=MUTED, first=True, tracking=100)
        tf = textbox(slide, rx + 1.62, ry, rw - 1.96, rh, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, value, size=13, color=TEXT, first=True, line_spacing=1.25)
        if i < n - 1:
            rect(slide, rx + 0.34, ry + rh - 0.01, rw - 0.68, 0.007, fill="DCE6DF")
        ry += rh

    footer(slide, foot_left(deck, s), "%02d" % s["index"])
    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 行动
def slide_cta(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide, dark=True)
    rect(slide, 0, 0, 0.14, SH, fill=AMBER)
    add_leaf(slide, 11.4, 2.0, 4.6, 1.6, -20, "FFFFFF", alpha=6)

    tf = textbox(slide, ML, 0.92, 9.0, 0.3)
    para(tf, s.get("kicker", ""), size=11, bold=True, color=AMBER_L, first=True, tracking=140)
    tf = textbox(slide, ML, 1.30, 10.6, 0.9)
    para(tf, s.get("title", ""), size=36, bold=True, color=CREAM, first=True, line_spacing=1.1)

    items = s.get("bullets") or []
    y = 2.70
    rh = min(0.86, (6.30 - y) / max(1, len(items)))
    for i, it in enumerate(items):
        yy = y + i * (rh + 0.02)
        oval(slide, ML + 0.02, yy + rh / 2 - 0.17, 0.34, 0.34, fill=AMBER)
        tf = textbox(slide, ML + 0.02, yy + rh / 2 - 0.17, 0.34, 0.34, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, str(i + 1), size=11.5, bold=True, color=INK, align=PP_ALIGN.CENTER, first=True)
        tf = textbox(slide, ML + 0.62, yy, 10.4, rh, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, it, size=16, color=ON_DARK, first=True, line_spacing=1.3)

    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 收尾
def slide_end(prs, deck, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide, dark=True)
    add_leaf(slide, 6.67, 3.05, 7.4, 2.6, -18, "FFFFFF", alpha=6)

    oval(slide, 5.42, 1.60, 2.5, 2.5, fill=None, line="254A40", line_w=1.2)
    add_leaf(slide, 6.67, 2.85, 1.30, 0.52, -28, AMBER_L, alpha=95)
    tf = textbox(slide, 1.2, 4.42, SW - 2.4, 0.7)
    para(tf, s.get("title", ""), size=30, bold=True, color=CREAM, align=PP_ALIGN.CENTER, first=True)
    if s.get("subtitle"):
        tf = textbox(slide, 1.6, 5.20, SW - 3.2, 0.5)
        para(tf, s["subtitle"], size=15, color=ON_DARK_MUTED, align=PP_ALIGN.CENTER, first=True)
    items = s.get("bullets") or []
    if items:
        tf = textbox(slide, 2.2, 5.86, SW - 4.4, 0.9)
        for i, it in enumerate(items):
            para(tf, it, size=11, color=ON_DARK_MUTED, align=PP_ALIGN.CENTER,
                 first=(i == 0), space_after=3, line_spacing=1.35)

    notes(slide, deck, s)
    return slide


# ---------------------------------------------------------------- 注册表
RENDERERS = {
    "cover": slide_cover,
    "points": slide_points,
    "cards": slide_cards,
    "hook": slide_cards,
    "tasting": slide_cards,
    "table": slide_table,
    "compare": slide_compare,
    "steps": slide_steps,
    "faq": slide_faq,
    "price": slide_price,
    "product": slide_product,
    "cta": slide_cta,
    "end": slide_end,
}


def resolve_renderer(s):
    """版式回退：按内容挑版式，避免未知类型（含历史存档里的旧类型）退化成空白页。"""
    if s.get("cards"):
        return slide_cards
    if s.get("faq"):
        return slide_faq
    if s.get("steps"):
        return slide_steps
    if s.get("table"):
        return slide_table
    if s.get("left") and s.get("right"):
        return slide_compare
    if s.get("price") or s.get("anchors"):
        return slide_price
    return slide_points


# ---------------------------------------------------------------- 入口
def build(deck, out_path):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    for s in deck.get("slides", []):
        fn = RENDERERS.get(s.get("type")) or resolve_renderer(s)
        try:
            fn(prs, deck, s)
        except Exception as e:  # 单页失败不影响整体导出
            print("[warn] slide %s (%s) 渲染失败: %s" % (s.get("id"), s.get("type"), e), file=sys.stderr)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    return len(prs.slides._sldIdLst)


def main():
    if len(sys.argv) < 3:
        print("usage: build_pptx.py <deck.json> <out.pptx>", file=sys.stderr)
        return 2
    with open(sys.argv[1], "r", encoding="utf-8") as f:
        deck = json.load(f)
    n = build(deck, sys.argv[2])
    print(json.dumps({"ok": True, "slides": n, "out": os.path.abspath(sys.argv[2])}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
